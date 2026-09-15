"""
Telugu Word Document Toolkit
Reusable helper library to create professional Microsoft Word (.docx) documents
with 100% correct Unicode Telugu fonts and formatting.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def format_telugu_run(run, text, font_name="Nirmala UI", size_pt=10, bold=False, color_rgb=None, italic=False):
    """
    Adds Telugu text with OpenXML Complex Script (w:cs) settings
    to ensure conjuncts, vattulu, and guninthalu render properly in Word.
    """
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name) # Key for Indic/Telugu rendering

def set_cell_background(cell, fill_hex):
    """Sets the background fill color of a table cell (e.g. '1E3A8A')."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_padding(cell, top=60, bottom=60, left=80, right=80):
    """Sets cell internal margins/padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def prevent_row_split(row):
    """Prevents a table row from awkwardly breaking across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    """Applies neat outer and inner grid borders to a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

class TeluguDocument:
    """Helper wrapper to quickly build Telugu documents with tables."""
    def __init__(self, orientation="landscape", margin_inches=0.4):
        self.doc = docx.Document()
        for s in self.doc.sections:
            if orientation == "landscape":
                s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
                s.page_width = Inches(11.69)
                s.page_height = Inches(8.27)
            else:
                s.orientation = docx.enum.section.WD_ORIENT.PORTRAIT
                s.page_width = Inches(8.27)
                s.page_height = Inches(11.69)
            s.top_margin = Inches(margin_inches)
            s.bottom_margin = Inches(margin_inches)
            s.left_margin = Inches(0.5)
            s.right_margin = Inches(0.5)

    def add_telugu_paragraph(self, text, size=11, bold=False, align="left", color=None):
        p = self.doc.add_paragraph()
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run()
        format_telugu_run(r, text, size_pt=size, bold=bold, color_rgb=color)
        return p

    def add_page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)
        r = p.add_run()
        r.add_break(WD_BREAK.PAGE)

    def save(self, filepath):
        self.doc.save(filepath)
