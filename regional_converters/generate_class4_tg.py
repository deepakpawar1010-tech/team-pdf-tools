import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=50, bottom=50, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def format_text_run(run, text, font_name="Nirmala UI", size_pt=9.5, bold=False, color_rgb=None, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def add_page_break_zero(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

def add_tg_header_bar(doc):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(5.35), Inches(5.35)]
    tbl.rows[0].cells[0].width = col_widths[0]
    tbl.rows[0].cells[1].width = col_widths[1]
    
    # Left pill: MICROSCHEDULE
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "DC2626")
    set_cell_margins(c0, top=50, bottom=50, left=140, right=140)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run()
    format_text_run(r0, "MICROSCHEDULE", font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    # Right pill: SEMESTER - I
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "DC2626")
    set_cell_margins(c1, top=50, bottom=50, left=140, right=140)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run()
    format_text_run(r1, "SEMESTER - I", font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

def add_tg_info_box(doc, lang_code="Telugu - FL", periods_week=6):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="6", val="single")
    
    col_widths = [Inches(2.5), Inches(3.7), Inches(2.3), Inches(2.2)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
    
    # Box 1: Subject
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "FEF2F2")
    set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run()
    format_text_run(r0, lang_code, font_name="Calibri", size_pt=13, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    # Box 2: Month & Dates
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "FFFBEB")
    set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1_1 = p1.add_run("September\n")
    format_text_run(r1_1, "September\n", font_name="Calibri", size_pt=13, bold=True, color_rgb=RGBColor(185, 28, 28))
    r1_2 = p1.add_run("31-08-2026 to 03-10-2026")
    format_text_run(r1_2, "31-08-2026 to 03-10-2026", font_name="Calibri", size_pt=12, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    # Box 3: Class
    c2 = tbl.rows[0].cells[2]
    set_cell_background(c2, "FEF2F2")
    set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run()
    format_text_run(r2, "Class-IV", font_name="Calibri", size_pt=14, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    # Box 4: Periods per week
    c3 = tbl.rows[0].cells[3]
    set_cell_background(c3, "E0F2FE")
    set_cell_margins(c3, top=60, bottom=60, left=100, right=100)
    c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(0)
    r3_1 = p3.add_run("No. of Periods per week  ")
    format_text_run(r3_1, "No. of Periods per week  ", font_name="Calibri", size_pt=10, bold=False, color_rgb=RGBColor(15, 23, 42))
    r3_2 = p3.add_run(str(periods_week))
    format_text_run(r3_2, str(periods_week), font_name="Calibri", size_pt=13, bold=True, color_rgb=RGBColor(15, 23, 42))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

def add_tg_footer(doc, page_no="2"):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(4.5), Inches(1.7), Inches(4.5)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
        
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "EA580C")
    set_cell_margins(c0, top=40, bottom=40, left=100, right=100)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run()
    format_text_run(r0, "NARAYANA GROUP OF SCHOOLS - INDIA", font_name="Arial", size_pt=10, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    c1 = tbl.rows[0].cells[1]
    set_cell_margins(c1, top=40, bottom=40, left=60, right=60)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run()
    format_text_run(r1, f"[ {page_no} ]", font_name="Arial", size_pt=11, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    c2 = tbl.rows[0].cells[2]
    set_cell_margins(c2, top=40, bottom=40, left=60, right=60)
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)

COL_WIDTHS_TG = [Inches(1.1), Inches(4.1), Inches(1.8), Inches(1.8), Inches(1.9)]

def init_tg_table(doc):
    tbl = doc.add_table(rows=0, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="4B5563", sz="4", val="single")
    
    row = tbl.add_row()
    set_row_cant_split(row)
    headers = ["PERIOD", "CLASSROOM TEACHING PHASE", "TEACHING /\\nAUDIOVISUAL\\nAIDS", "ENRICHMENT\\nACTIVITIES", "HOME TASK"]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.width = COL_WIDTHS_TG[i]
        set_cell_background(cell, "2E7D32") # Narayana Green
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        lines = h.split("\\n")
        for li, l in enumerate(lines):
            if li > 0:
                p.add_run("\\n")
            r = p.add_run()
            format_text_run(r, l, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(255, 255, 255))
    return tbl

def add_tg_banner(tbl, text, b_type="date"):
    row = tbl.add_row()
    set_row_cant_split(row)
    cell = row.cells[0]
    for c in row.cells[1:]:
        cell.merge(c)
        
    set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    if b_type == "date":
        set_cell_background(cell, "A9DFBF") # Soft light green
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        format_text_run(r, text, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(20, 83, 45))
    elif b_type == "note_red":
        set_cell_background(cell, "FFFFFF")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lines = text.split("\\n")
        for li, l in enumerate(lines):
            if li > 0:
                p.add_run("\\n")
            r = p.add_run()
            format_text_run(r, l, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(220, 38, 38))
    return row

def add_tg_data_row(tbl, period="", classroom="", teaching="", enrichment="", hometask="",
                    classroom_bold=False, span_all_after_classroom=None):
    row = tbl.add_row()
    set_row_cant_split(row)
    
    if span_all_after_classroom:
        # cols 2, 3, 4 are merged with span_all_after_classroom text
        c2 = row.cells[2]
        for c in row.cells[3:]:
            c2.merge(c)
        values = [str(period), str(classroom), str(span_all_after_classroom)]
        col_indices = [0, 1, 2]
    else:
        values = [str(period), str(classroom), str(teaching), str(enrichment), str(hometask)]
        col_indices = [0, 1, 2, 3, 4]
        
    for idx_val, i in enumerate(col_indices):
        cell = row.cells[i]
        if not span_all_after_classroom or i < 2:
            cell.width = COL_WIDTHS_TG[i]
        set_cell_margins(cell, top=45, bottom=45, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        val = values[idx_val]
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                r = p.add_run()
                format_text_run(r, val, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(15, 23, 42))
        elif i == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                lines = val.split("\\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\\n")
                    r = p.add_run()
                    # Check if Joyful saturday or lesson title
                    is_red = "JOYFUL SATURDAY" in l
                    col = RGBColor(220, 38, 38) if is_red else RGBColor(15, 23, 42)
                    format_text_run(r, l, font_name="Nirmala UI", size_pt=9.5, bold=(classroom_bold or is_red), color_rgb=col)
        elif span_all_after_classroom:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                r = p.add_run()
                is_red = "సామాగ్రి" in val
                col = RGBColor(220, 38, 38) if is_red else RGBColor(15, 23, 42)
                format_text_run(r, val, font_name="Nirmala UI", size_pt=10, bold=True, color_rgb=col)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                lines = val.split("\\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\\n")
                    r = p.add_run()
                    format_text_run(r, l, font_name="Nirmala UI", size_pt=9, bold=False, color_rgb=RGBColor(15, 23, 42))
    return row

def generate_tg_fl_doc(output_path):
    """Generates CLASS-IV-TG-MCS-SEPTEMBER-pages-2-3 (Telugu - First Language)"""
    print("Generating Telugu - FL (Pages 2-3)...")
    doc = docx.Document()
    for s in doc.sections:
        s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        s.page_width = Inches(11.69)
        s.page_height = Inches(8.27)
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)
        
    # ==================== PAGE 1 (Page 2 of MCS) ====================
    add_tg_header_bar(doc)
    add_tg_info_box(doc, lang_code="Telugu - FL", periods_week=6)
    t1 = init_tg_table(doc)
    
    add_tg_banner(t1, "31-08-2026 to 12-09-2026", "date")
    add_tg_data_row(t1, period="Period 1", 
                    classroom="పాఠం 3. ఊరు  TB Pg No. 65 లో ఉన్న అభ్యాస పత్రాన్ని బోధించుట, రాయించుట, చదివించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="-")
    add_tg_data_row(t1, period="Period 2", 
                    classroom="పాఠం 3. ఊరు  TB Pg No. 66, 67 లో ఉన్న అభ్యాస పత్రాన్ని బోధించుట, రాయించుట, చదివించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="TB Pg No.66\\n(ఉ) (bit)")
                    
    note_p1 = ("Note: Formative Assessment II Revision (01-09-2026 to 03-09-2026)\\n"
               "         Formative Assessment II (07-09-2026 to 10-09-2026)\\n"
               "         Janmashtami (04-09-2026)\\n"
               "         Teachers’ Day/Buffer (05-09-2026)\\n"
               "         Second Saturday (12-09-2026)")
    add_tg_banner(t1, note_p1, "note_red")
    
    add_tg_banner(t1, "14-09-2026 to 19-09-2026", "date")
    add_tg_data_row(t1, period="Period 1", classroom="పాఠం 3. ఊరు", span_all_after_classroom="REVISION")
    add_tg_data_row(t1, period="Period 2", classroom="పాఠం 3. ఊరు", span_all_after_classroom="SLIPTEST")
    add_tg_data_row(t1, period="Period 3", 
                    classroom="సంస్కృతం (वर्णमाला) (వర్ణమాల) TB Pg No.68 లో ఉన్న అభ్యాసాలను బోధించుట, రాయించుట, చదివించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="సంస్కృత పదాలకు\\nతెలుగులో పేర్లు\\nచెప్పించుట.", hometask="-")
                    
    add_tg_footer(doc, page_no="2")
    
    # ==================== PAGE 2 (Page 3 of MCS) ====================
    add_page_break_zero(doc)
    add_tg_header_bar(doc)
    t2 = init_tg_table(doc)
    
    add_tg_data_row(t2, period="Period 4", 
                    classroom="పాఠం 4. పర్యావరణం-పరిరక్షణ  TB Pg No.69 లో ఉన్న చిత్రాన్ని చూపించుట, చిత్రం గురించి విద్యార్థులచే మాట్లాడించుట, TB Pg No.70 లో ఉన్న పాఠాన్ని బోధించుట.",
                    teaching="పాఠ్యపుస్తకం\\nపాఠానికి సంబంధించిన\\nబోధనోపకరణాలను\\nచూపించుట.", 
                    enrichment="పర్యావరణం గురించి\\nవిద్యార్థులచే\\nచెప్పించుట.", hometask="-")
    add_tg_data_row(t2, period="Period 5", classroom="", span_all_after_classroom="AV")
    
    add_tg_banner(t2, "Note: Ganesha Chaturthi (14-09-2026)", "note_red")
    add_tg_banner(t2, "21-09-2026 to 26-09-2026", "date")
    
    add_tg_data_row(t2, period="Period 1", 
                    classroom="పాఠం 4. పర్యావరణం-పరిరక్షణ  TB Pg No.71, 72 లో ఉన్న పాఠాన్ని బోధించుట.",
                    teaching="పాఠ్యపుస్తకం\\nపాఠానికి సంబంధించిన\\nబోధనోపకరణాలను\\nచూపించుట. AV", 
                    enrichment="-", hometask="పాఠాన్ని\\nచదువుకొని\\nవచ్చుట.")
    add_tg_data_row(t2, period="Period 2", 
                    classroom="పాఠం 4. పర్యావరణం-పరిరక్షణ  TB Pg No.70, 71, 72 లో ఉన్న పాఠాన్ని విద్యార్థులచే సంభాషణ రూపంలో చెప్పించుట.",
                    teaching="పాఠ్యపుస్తకం\\nపాఠానికి సంబంధించిన\\nబోధనోపకరణాలను\\nచూపించుట. AV", 
                    enrichment="విద్యార్థులచే సంభాషణ\\nరూపంలో చెప్పించుట.", hometask="-")
    add_tg_data_row(t2, period="Period 3", 
                    classroom="పాఠం 4. పర్యావరణం-పరిరక్షణ  TB Pg No.73\\nI, II రోమన్ లలో ఉన్న అభ్యాసాలను బోధించుట, రాయించుట, చదివించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="-")
    add_tg_data_row(t2, period="Period 4", 
                    classroom="పాఠం 4. పర్యావరణం-పరిరక్షణ  TB Pg No.74\\nII, III రోమన్ లలో ఉన్న అభ్యాసాలను బోధించుట, రాయించుట, చదివించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="ప్రశ్న జవాబులను\\nచదువుకొని వచ్చుట.")
                    
    add_tg_footer(doc, page_no="3")
    
    doc.save(output_path)
    print(f"Saved: {output_path}")

def generate_tg_sl_doc(output_path):
    """Generates CLASS-IV-TG-MCS-SEPTEMBER-pages-35-36 (Telugu - Second Language)"""
    print("Generating Telugu - SL (Pages 35-36)...")
    doc = docx.Document()
    for s in doc.sections:
        s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        s.page_width = Inches(11.69)
        s.page_height = Inches(8.27)
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)
        
    # ==================== PAGE 1 (Page 35 of MCS) ====================
    add_tg_header_bar(doc)
    add_tg_info_box(doc, lang_code="Telugu - SL", periods_week=4)
    t1 = init_tg_table(doc)
    
    add_tg_banner(t1, "31-08-2026 to 12-09-2026", "date")
    add_tg_data_row(t1, period="Period 1", classroom="పాఠం-4. కాకమ్మ!", span_all_after_classroom="SLIPTEST")
    
    note_p35 = ("Note: Formative Assessment II Revision (01-09-2026 to 03-09-2026)\\n"
                "         Formative Assessment II (07-09-2026 to 10-09-2026)\\n"
                "         Janmashtami (04-09-2026)\\n"
                "         Teachers’ Day/Buffer (05-09-2026)\\n"
                "         Second Saturday (12-09-2026)")
    add_tg_banner(t1, note_p35, "note_red")
    
    add_tg_banner(t1, "14-09-2026 to 19-09-2026", "date")
    add_tg_data_row(t1, period="Period 1", 
                    classroom="(వినండి-పాడండి.) ఇదే మా గుడి\\nTB Pg No.57 లో ఉన్న గేయాన్ని బోధించుట, అభినయిస్తూ పాడించుట.",
                    teaching="పాఠానికి సంబంధించిన\\nబోధనోపకరణాలను\\nచూపించుట. AV", 
                    enrichment="గేయాన్ని\\nఅభినయిస్తూ\\nపాడుతారు.", hometask="-")
    add_tg_data_row(t1, period="Period 2", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య TB Pg No.58, 59 లో ఉన్న గేయాలను, ద్విత్వాక్షరాలను, పదాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="గేయాలను\\nఅభినయిస్తూ\\nపాడుతారు.", hometask="-")
    add_tg_data_row(t1, period="Period 3", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య TB Pg No.60 లో ఉన్న గేయాలను, ద్విత్వాక్షరాలను, పదాలను, TB Pg No.61 లో ఉన్న అభ్యాసాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం\\nపాఠానికి సంబంధించిన\\nబోధనోపకరణాలను\\nచూపించుట. AV", 
                    enrichment="TB Pg No.61\\nII రోమన్ (ఆ) (bit)", 
                    hometask="TB Pg No.60\\nలో ఉన్న గేయాలను\\nచదువుకొని వచ్చుట.")
                    
    add_tg_footer(doc, page_no="35")
    
    # ==================== PAGE 2 (Page 36 of MCS) ====================
    add_page_break_zero(doc)
    add_tg_header_bar(doc)
    t2 = init_tg_table(doc)
    
    add_tg_data_row(t2, period="Period 4", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య  TB Pg No.62, 63 లో ఉన్న అభ్యాసాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం", 
                    enrichment="TB Pg No.63\\nIII రోమన్ (ఆ) (bit)", 
                    hometask="TB Pg No.62\\nIII రోమన్ (ఆ) (bit)\\n“ క్క ” గుణింతం\\nరాసుకొని వచ్చుట.")
                    
    add_tg_banner(t2, "Note: Ganesha Chaturthi (14-09-2026)", "note_red")
    add_tg_banner(t2, "21-09-2026 to 26-09-2026", "date")
    
    add_tg_data_row(t2, period="Period 1", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య  TB Pg No.64, 65 లో ఉన్న అభ్యాసాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="TB Pg No.65\\nV రోమన్ (ఆ) (bit)", hometask="-")
    add_tg_data_row(t2, period="Period 2", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య  TB Pg No.66, 67 లో ఉన్న అభ్యాసాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="-")
    add_tg_data_row(t2, period="Period 3", 
                    classroom="పాఠం 5. నక్క-కన్నయ్య  TB Pg No. 68, 69 లో ఉన్న అభ్యాస పత్రాలను బోధించుట, చదివించుట, రాయించుట.",
                    teaching="పాఠ్యపుస్తకం", enrichment="-", hometask="TB Pg No.68\\n(ఇ) (bit)")
    add_tg_data_row(t2, period="Period 4", 
                    classroom="JOYFUL SATURDAY :-\\nకన్నయ్య బొమ్మను అందంగా గీసి, రంగులు వేసి, పేరు రాయండి.",
                    span_all_after_classroom="సామాగ్రి :- చార్టు, పెన్సిల్, కలర్ పెన్సిల్స్")
                    
    add_tg_banner(t2, "Note: Joyful Saturday (26-09-2026)", "note_red")
    add_tg_banner(t2, "28-09-2026 to 03-10-2026", "date")
    
    add_tg_data_row(t2, period="Period 1", classroom="పాఠం 5. నక్క-కన్నయ్య", span_all_after_classroom="REVISION")
    add_tg_data_row(t2, period="Period 2", classroom="పాఠం 5. నక్క-కన్నయ్య", span_all_after_classroom="SLIPTEST")
    
    add_tg_footer(doc, page_no="36")
    
    doc.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\Deepak Pawar\Desktop\Automation Projects\Telugu_PDF_to_Word\output_docs"
    split_dir = r"C:\Users\Deepak Pawar\Downloads\CLASS-IV-TG-MCS-SEPTEMBER-split-files"
    
    f1_out1 = os.path.join(out_dir, "CLASS-IV-TG-MCS-SEPTEMBER-pages-2-3.docx")
    f1_out2 = os.path.join(split_dir, "CLASS-IV-TG-MCS-SEPTEMBER-pages-2-3.docx")
    generate_tg_fl_doc(f1_out1)
    import shutil
    shutil.copy(f1_out1, f1_out2)
    
    f2_out1 = os.path.join(out_dir, "CLASS-IV-TG-MCS-SEPTEMBER-pages-35-36.docx")
    f2_out2 = os.path.join(split_dir, "CLASS-IV-TG-MCS-SEPTEMBER-pages-35-36.docx")
    generate_tg_sl_doc(f2_out1)
    shutil.copy(f2_out1, f2_out2)
    
    print("Completed both documents!")
