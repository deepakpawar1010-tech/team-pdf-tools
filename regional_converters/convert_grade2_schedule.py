import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_borders(table, color="7F9DB9", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def format_run(run, font_name="Nirmala UI", size_pt=9.5, bold=False, color_rgb=None, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    # Word Complex Script font for Telugu Unicode support (preserves conjuncts and vattulu)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def add_header_bar(doc):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(3.2), Inches(4.3), Inches(3.2)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
    
    # Left pill: Micro schedule
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "DC2626")
    set_cell_margins(c0, top=60, bottom=60, left=140, right=140)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run("Micro schedule")
    format_run(r0, font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    # Center pill: AY-2026 (Term - 1)
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "EA580C")
    set_cell_margins(c1, top=60, bottom=60, left=140, right=140)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run("AY-2026 (Term - 1)")
    format_run(r1, font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    # Right pill: Andhra Pradesh
    c2 = tbl.rows[0].cells[2]
    set_cell_background(c2, "DC2626")
    set_cell_margins(c2, top=60, bottom=60, left=140, right=140)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Andhra Pradesh")
    format_run(r2, font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

def add_info_box(doc):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="6", val="single")
    
    col_widths = [Inches(2.5), Inches(2.7), Inches(2.3), Inches(3.2)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
    
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "FEF2F2")
    set_cell_margins(c0, top=50, bottom=50, left=100, right=100)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0_1 = p0.add_run("No.of Periods\n")
    format_run(r0_1, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(185, 28, 28))
    r0_2 = p0.add_run("Per Week: 4")
    format_run(r0_2, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "FFFBEB")
    set_cell_margins(c1, top=50, bottom=50, left=100, right=100)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1_1 = p1.add_run("April to October\n")
    format_run(r1_1, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(180, 83, 9))
    r1_2 = p1.add_run("2026-27")
    format_run(r1_2, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(180, 83, 9))
    
    c2 = tbl.rows[0].cells[2]
    set_cell_background(c2, "EFF6FF")
    set_cell_margins(c2, top=50, bottom=50, left=100, right=100)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Grade: II")
    format_run(r2, font_name="Calibri", size_pt=12, bold=True, color_rgb=RGBColor(29, 78, 216))
    
    c3 = tbl.rows[0].cells[3]
    set_cell_background(c3, "F5F3FF")
    set_cell_margins(c3, top=50, bottom=50, left=100, right=100)
    c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(0)
    r3_1 = p3.add_run("TELUGU\n")
    format_run(r3_1, font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(67, 56, 202))
    r3_2 = p3.add_run("Subject: Third Language")
    format_run(r3_2, font_name="Calibri", size_pt=10.5, bold=True, color_rgb=RGBColor(185, 28, 28))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

COL_WIDTHS = [Inches(1.0), Inches(2.2), Inches(4.2), Inches(1.8), Inches(1.5)]

def init_schedule_table(doc):
    tbl = doc.add_table(rows=0, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="94A3B8", sz="4", val="single")
    
    row = tbl.add_row()
    set_row_cant_split(row)
    headers = ["PERIOD", "CHAPTER/LESSON", "CONTENT TO BE TAUGHT", "TLM", "PRACTICE WORK"]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.width = COL_WIDTHS[i]
        set_cell_background(cell, "1E3A8A") # Navy Blue
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        format_run(r, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(255, 255, 255))
    return tbl

def add_banner(tbl, text, b_type="date"):
    row = tbl.add_row()
    set_row_cant_split(row)
    cell = row.cells[0]
    for c in row.cells[1:]:
        cell.merge(c)
    
    set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    if b_type == "date":
        set_cell_background(cell, "BFDBFE")
        r = p.add_run(text)
        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(15, 23, 42))
    elif b_type == "holiday_red":
        set_cell_background(cell, "FEE2E2")
        r = p.add_run(text)
        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(185, 28, 28))
    elif b_type == "event_purple":
        set_cell_background(cell, "F3E8FF")
        r = p.add_run(text)
        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(107, 33, 168))
    elif b_type == "event_green":
        set_cell_background(cell, "DCFCE7")
        r = p.add_run(text)
        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(21, 128, 61))
    elif b_type == "blue_banner":
        set_cell_background(cell, "E0E7FF")
        r = p.add_run(text)
        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(30, 58, 138))
    return row

def add_data_row(tbl, period="", chapter="", content="", tlm="", practice="", 
                 chapter_color="green", content_color="black"):
    row = tbl.add_row()
    set_row_cant_split(row)
    vals = [str(period), str(chapter), str(content), str(tlm), str(practice)]
    
    for i in range(5):
        cell = row.cells[i]
        cell.width = COL_WIDTHS[i]
        set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        val = vals[i]
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                r = p.add_run(val)
                format_run(r, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(15, 23, 42))
        elif i == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                r = p.add_run(val)
                col = RGBColor(16, 124, 65) if chapter_color == "green" else RGBColor(0, 0, 0)
                format_run(r, font_name="Nirmala UI", size_pt=9.5, bold=True, color_rgb=col)
        elif i == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                lines = val.split("\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\n")
                    r = p.add_run(l)
                    if content_color == "green":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(16, 124, 65))
                    elif content_color == "red" or "లఘు పరీక్ష" in l or "పునశ్చరణ" in l:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        format_run(r, font_name="Nirmala UI", size_pt=10, bold=True, color_rgb=RGBColor(185, 28, 28))
                    else:
                        format_run(r, font_name="Nirmala UI", size_pt=9, bold=False, color_rgb=RGBColor(15, 23, 42))
        elif i == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                lines = val.split("\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\n")
                    r = p.add_run(l)
                    if l.strip() == "AV MODULES":
                        format_run(r, font_name="Calibri", size_pt=9, bold=True, color_rgb=RGBColor(30, 58, 138))
                    else:
                        format_run(r, font_name="Nirmala UI", size_pt=8.5, bold=False, color_rgb=RGBColor(30, 41, 59))
        elif i == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                lines = val.split("\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\n")
                    r = p.add_run(l)
                    format_run(r, font_name="Nirmala UI", size_pt=8.5, bold=False, color_rgb=RGBColor(30, 41, 59))
    return row

def add_page_break_zero(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

def convert():
    print("=========================================================")
    print(" Telugu Micro Schedule Word (.docx) Generator")
    print(" Target: Grade II Telugu Third Language (Term 1)")
    print("=========================================================")
    
    doc = docx.Document()
    
    for s in doc.sections:
        s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        s.page_width = Inches(11.69)
        s.page_height = Inches(8.27)
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)
    
    # Page 1
    print("[1/8] Building Page 1...")
    add_header_bar(doc)
    add_info_box(doc)
    t1 = init_schedule_table(doc)
    add_banner(t1, "01-04-2026 TO 04-04-2026", "date")
    add_data_row(t1, period="1", content="RE - OPEN", content_color="green")
    add_data_row(t1, period="2", chapter="n-Fundamentals", 
                 content="1. బాణం గుర్తుల ఆధారంగా రాయుట", 
                 tlm="నల్లబల్ల వర్ణమాల చార్టు", practice="గుర్తులు రాయించుట")
    add_banner(t1, "01-04-2026 : BUFFER", "event_purple")
    add_banner(t1, "03-04-2026 : GOOD FRIDAY", "holiday_red")
    add_banner(t1, "06-04-2026 TO 11-04-2026", "date")
    r1 = add_data_row(t1, period="1", chapter="n-Fundamentals", 
                      content="2. వర్ణమాల - అచ్చులు (అ నుండి ౠ వరకు)", 
                      practice="అచ్చులు రాయించుట")
    r2 = add_data_row(t1, period="2", 
                      content="3. వర్ణమాల - అచ్చులు (ఎ నుండి అః వరకు)", 
                      practice="అచ్చులు రాయించుట")
    r3 = add_data_row(t1, period="3", 
                      content="4. వర్ణమాల - అచ్చులు (అ నుండి అః వరకు)", 
                      tlm="నల్లబల్ల వర్ణమాల చార్టు", practice="అచ్చులు రాయించుట")
    r4 = add_data_row(t1, period="4", 
                      content="5. హల్లులు (క నుండి ఞ వరకు)", 
                      tlm="నల్లబల్ల వర్ణమాల చార్టు", practice="హల్లులు రాయించుట")
    r1.cells[1].merge(r2.cells[1]).merge(r3.cells[1]).merge(r4.cells[1])
    add_banner(t1, "11-04-2026 : SECOND SATURDAY", "holiday_red")
    add_banner(t1, "13-04-2026 TO 18-04-2026", "date")
    r5 = add_data_row(t1, period="1", chapter="n-Fundamentals", 
                      content="6. హల్లులు (ట నుండి న వరకు)")
    r6 = add_data_row(t1, period="2", 
                      content="7. హల్లులు (ప నుండి వ వరకు)", 
                      practice="హల్లులు రాయించుట")
    r7 = add_data_row(t1, period="3", 
                      content="8. హల్లులు (శ నుండి ఱ వరకు)", 
                      tlm="నల్లబల్ల వర్ణమాల చార్టు", practice="హల్లులు రాయించుట")
    r8 = add_data_row(t1, period="4", 
                      content="9. వర్ణమాల (అ నుండి ఱ వరకు)", 
                      practice="హల్లులు రాయించుట")
    r5.cells[1].merge(r6.cells[1]).merge(r7.cells[1]).merge(r8.cells[1])
    add_banner(t1, "14-04-2026 : AMBEDKAR JAYANTI", "holiday_red")
    
    # Page 2
    print("[2/8] Building Page 2...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t2 = init_schedule_table(doc)
    add_banner(t2, "20-04-2026 TO 25-04-2026", "date")
    add_data_row(t2, period="1", chapter="n-Fundamentals", 
                 content="10. వర్ణమాల (అ నుండి ఱ వరకు)", 
                 tlm="నల్లబల్ల వర్ణమాల చార్టు")
    add_data_row(t2, period="2", chapter="అక్షరమాల - అచ్చులు", 
                 content="అచ్చులు పలకండి")
    add_data_row(t2, period="3", chapter="అక్షరమాల - హల్లులు", 
                 content="హల్లులు పలకండి")
    add_data_row(t2, period="4", chapter="రాయడం ఇలా", 
                 content="బాణం గుర్తుల ఆధారంగా రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 10, 11")
    add_banner(t2, "27-04-2026 TO 11-06-2026 : SUMMER VACATION", "holiday_red")
    add_banner(t2, "12-06-2026 & 13-06-2026\nRE OPENING DAY", "event_green")
    add_banner(t2, "13-06-2026 : SECOND SATURDAY", "holiday_red")
    add_banner(t2, "15-06-2026 TO 20-06-2026", "date")
    add_data_row(t2, period="1", chapter="అక్షరమాల - అచ్చులు\n(పునశ్చరణ)", 
                 content="అచ్చులు పలకండి", tlm="నల్లబల్ల")
    add_data_row(t2, period="2", chapter="అక్షరమాల - హల్లులు\n(పునశ్చరణ)", 
                 content="హల్లులు పలకండి", tlm="పాఠ్యపుస్తకం,\nవర్ణమాల చార్టు, నల్ల బల్ల")
    add_data_row(t2, period="3", chapter="రాయడం ఇలా\n(పునశ్చరణ)", 
                 content="బాణం గుర్తుల ఆధారంగా రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 10, 11", 
                 tlm="పాఠ్యపుస్తకం, నల్ల బల్ల")
    add_banner(t2, "20-06-2026 : NO BAG DAY – STEAM & CCA", "event_purple")
    add_banner(t2, "22-06-2026 TO 27-06-2026", "date")
    r9 = add_data_row(t2, period="1", chapter="1. అమ్మ - ఆవు (అ, ఆ)", 
                      content="బొమ్మలను చూడండి-వాటి పేర్లు చదవండి, చెప్పండి\nపాఠ్యబోధన పేజీ నెం. 12", 
                      tlm="పాఠ్యపుస్తకం,\nవర్ణమాల చార్టు")
    r10 = add_data_row(t2, period="2", 
                       content="చుక్కలను కలుపుతూ అక్షరాలు రాయుట", 
                       tlm="పాఠ్యపుస్తకం\nAV MODULES")
    r11 = add_data_row(t2, period="3", 
                       content="బొమ్మలను చూడండి-వాటి పేర్లు చదవండి, చెప్పండి\nపాఠ్యబోధన పేజీ నెం. 13", 
                       tlm="పాఠ్యపుస్తకం,\nవర్ణమాల చార్టు")
    r9.cells[1].merge(r10.cells[1]).merge(r11.cells[1])
    
    # Page 3
    print("[3/8] Building Page 3...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t3 = init_schedule_table(doc)
    add_data_row(t3, period="4", chapter="1. అమ్మ - ఆవు (అ, ఆ)", 
                 content="అభ్యాసం చేయుట")
    add_banner(t3, "26-06-2026 : MUHARRAM", "holiday_red")
    add_banner(t3, "27-06-2026 : NO BAG DAY/BB/CCA/CLUB", "event_purple")
    add_banner(t3, "29-06-2026 TO 04-07-2026", "date")
    r12 = add_data_row(t3, period="1", chapter="1. అమ్మ - ఆవు (అ, ఆ)", 
                       content="అభ్యాసం :\nఅభ్యాసాలు చేయించుట", 
                       tlm="పాఠ్యపుస్తకం,\nవర్ణమాల చార్టు", 
                       practice="తరగతి పని అ-ఆ\nఅక్షరాలు రాయుట")
    r13 = add_data_row(t3, period="2")
    r14 = add_data_row(t3, period="3", tlm="పాఠ్యపుస్తకం")
    r12.cells[1].merge(r13.cells[1]).merge(r14.cells[1])
    r12.cells[2].merge(r13.cells[2]).merge(r14.cells[2])
    r12.cells[4].merge(r13.cells[4]).merge(r14.cells[4])
    add_data_row(t3, period="4", chapter="1. అమ్మ - ఆవు (అ, ఆ)", 
                 content="లఘు పరీక్ష", content_color="red")
    add_banner(t3, "29-06-2026 & 30-06-2026 : PTM", "event_purple")
    add_banner(t3, "04-07-2026 : NO BAG DAY - STEAM & CCA", "event_purple")
    add_banner(t3, "06-07-2026 TO 11-07-2026", "date")
    r15 = add_data_row(t3, period="1", chapter="2. ఇల్లు - ఈగ (ఇ, ఈ)", 
                       content="బొమ్మలను చూడండి-వాటి పేర్లు చదవండి, చెప్పండి\nపాఠ్యబోధన పేజీ నెం. 15\nచుక్కలను కలుపుతూ అక్షరాలు రాయుట", 
                       tlm="పాఠ్యపుస్తకం,\nవర్ణమాల చార్టు")
    r16 = add_data_row(t3, period="2", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట,\nపాఠ్యపుస్తకం పేజీ నెం. 15", 
                       tlm="పాఠ్యపుస్తకం\nAV MODULES")
    r17 = add_data_row(t3, period="3", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట,\nపాఠ్యపుస్తకం పేజీ నెం. 16", 
                       tlm="పాఠ్యపుస్తకం")
    r18 = add_data_row(t3, period="4", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 16")
    r15.cells[1].merge(r16.cells[1]).merge(r17.cells[1]).merge(r18.cells[1])
    add_banner(t3, "11-07-2026 : SECOND SATURDAY", "holiday_red")
    
    # Page 4
    print("[4/8] Building Page 4...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t4 = init_schedule_table(doc)
    add_banner(t4, "13-07-2026 TO 18-07-2026", "date")
    r19 = add_data_row(t4, period="1", chapter="2. ఇల్లు - ఈగ (ఇ, ఈ)", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 17")
    r20 = add_data_row(t4, period="2", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 18", 
                       tlm="వర్ణమాల చార్టు,\nపాఠ్యపుస్తకం")
    r21 = add_data_row(t4, period="3", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 19")
    r19.cells[1].merge(r20.cells[1]).merge(r21.cells[1])
    add_data_row(t4, period="4", chapter="2. ఇల్లు - ఈగ (ఇ, ఈ)", 
                 content="లఘు పరీక్ష", content_color="red")
    add_banner(t4, "20-07-2026 TO 25-07-2026", "date")
    r22 = add_data_row(t4, period="1", chapter="3. ఉడత - ఊయల\n(ఉ, ఊ)", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట,\nపాఠ్యపుస్తకం పేజీ నెం. 20", 
                       tlm="పాఠ్యపుస్తకం, నల్లబల్ల")
    r23 = add_data_row(t4, period="2", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 21", 
                       tlm="పాఠ్యపుస్తకం\nAV MODULES")
    r24 = add_data_row(t4, period="3", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 22", 
                       tlm="పాఠ్యపుస్తకం, నల్లబల్ల,\nనోటు పుస్తకం")
    r22.cells[1].merge(r23.cells[1]).merge(r24.cells[1])
    add_banner(t4, "25-07-2026 : NO BAG DAY/BB/CCA/CLUB", "event_purple")
    add_banner(t4, "27-07-2026 TO 01-08-2026", "date")
    r25 = add_data_row(t4, period="1", chapter="3. ఉడత - ఊయల\n(ఉ, ఊ)", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 23", 
                       tlm="వర్ణమాల చార్టు,")
    r26 = add_data_row(t4, period="2", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 24", 
                       tlm="పాఠ్యపుస్తకం")
    r25.cells[1].merge(r26.cells[1])
    
    # Page 5
    print("[5/8] Building Page 5...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t5 = init_schedule_table(doc)
    add_data_row(t5, period="3", chapter="3. ఉడత - ఊయల\n(ఉ, ఊ)", 
                 content="లఘు పరీక్ష", content_color="red")
    add_banner(t5, "29-07-2026 TO 31-07-2026 : PTM", "event_purple")
    add_banner(t5, "03-08-2026 TO 08-08-2026", "date")
    add_data_row(t5, period="1", chapter="చిన్నారి పాపలం", 
                 content="చదువు - ఆనందించు")
    r27 = add_data_row(t5, period="2", chapter="4. ఋషి - ఐదు\n(ఋ, ఎ, ఏ, ఐ)", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట,\nపాఠ్యపుస్తకం పేజీ నెం. 26", 
                       tlm="పాఠ్యపుస్తకం,\nనల్లబల్ల")
    r28 = add_data_row(t5, period="3", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 27", 
                       tlm="పాఠ్యపుస్తకం\nAV MODULES")
    r29 = add_data_row(t5, period="4", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 28", 
                       tlm="పాఠ్యపుస్తకం, నల్లబల్ల,\nనోటు పుస్తకం")
    r27.cells[1].merge(r28.cells[1]).merge(r29.cells[1])
    add_banner(t5, "08-08-2026 : SECOND SATURDAY", "holiday_red")
    add_banner(t5, "10-08-2026 TO 15-08-2026", "date")
    r30 = add_data_row(t5, period="1", chapter="4. ఋషి - ఐదు\n(ఋ, ఎ, ఏ, ఐ)", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 29")
    r31 = add_data_row(t5, period="2", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట\nపాఠ్యపుస్తకం పేజీ నెం. 30", 
                       tlm="వర్ణమాల చార్టు")
    r32 = add_data_row(t5, period="3", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 31")
    r30.cells[1].merge(r31.cells[1]).merge(r32.cells[1])
    
    # Page 6
    print("[6/8] Building Page 6...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t6 = init_schedule_table(doc)
    add_data_row(t6, period="4", chapter="4. ఋషి - ఐదు\n(ఋ, ఎ, ఏ, ఐ)", 
                 content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 32", 
                 tlm="పాఠ్యపుస్తకం")
    add_banner(t6, "15-08-2026 : INDEPENDENCE DAY", "holiday_red")
    add_banner(t6, "17-08-2026 TO 22-08-2026", "date")
    add_data_row(t6, period="1", chapter="4. ఋషి - ఐదు\n(ఋ, ఎ, ఏ, ఐ)", 
                 content="లఘు పరీక్ష", content_color="red")
    r33 = add_data_row(t6, period="2", chapter="5. ఒంటె-అంతఃపురం\n(ఒ, ఓ, ఔ, అం, అః)", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట, పాఠ్యపుస్తకం\nపేజీ నెం. 33", 
                       tlm="పాఠ్యపుస్తకం, నల్లబల్ల")
    r34 = add_data_row(t6, period="3", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట, పాఠ్యపుస్తకం\nపేజీ నెం. 34", 
                       tlm="పాఠ్యపుస్తకం\nAV MODULES")
    r33.cells[1].merge(r34.cells[1])
    add_banner(t6, "21-08-2026 : VARALAKSHMI VRATHAM", "holiday_red")
    add_banner(t6, "22-08-2026 : NO BAG DAY-STEAM&CCA", "event_purple")
    add_banner(t6, "24-08-2026 TO 29-08-2026", "date")
    r35 = add_data_row(t6, period="1", chapter="5. ఒంటె-అంతఃపురం\n(ఒ, ఓ, ఔ, అం, అః)", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట, పాఠ్యపుస్తకం\nపేజీ నెం. 35", 
                       tlm="పాఠ్యపుస్తకం, నల్లబల్ల")
    r36 = add_data_row(t6, period="2", 
                       content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట, పాఠ్యపుస్తకం\nపేజీ నెం. 36")
    r35.cells[1].merge(r36.cells[1])
    
    # Page 7
    print("[7/8] Building Page 7...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t7 = init_schedule_table(doc)
    add_data_row(t7, period="3", chapter="5. ఒంటె-అంతఃపురం\n(ఒ, ఓ, ఔ, అం, అః)", 
                 content="బొమ్మలు చూసి అక్షరాలు చదువుట, చుక్కలను\nకలుపుతూ అక్షరాలు రాయుట, పాఠ్యపుస్తకం\nపేజీ నెం. 37")
    add_banner(t7, "25-08-2026 TO 28-08-2026 : PTM", "event_purple")
    add_banner(t7, "26-08-2026 : EID MILAD UN NABI", "holiday_red")
    add_banner(t7, "29-08-2026 : NO BAG DAY/BB/CCA/CLUB", "event_purple")
    add_banner(t7, "31-08-2026 TO 05-09-2026", "date")
    r37 = add_data_row(t7, period="1", chapter="5. ఒంటె-అంతఃపురం\n(ఒ, ఓ, ఔ, అం, అః)", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 38", 
                       tlm="పాఠ్యపుస్తకం,\nనల్ల బల్ల, నోటు పుస్తకం")
    r38 = add_data_row(t7, period="2", 
                       content="అభ్యాస పత్రం\nపాఠ్యపుస్తకం పేజీ నెం. 39")
    r37.cells[1].merge(r38.cells[1])
    add_data_row(t7, period="3", chapter="5. ఒంటె-అంతఃపురం\n(ఒ, ఓ, ఔ, అం, అః)", 
                 content="లఘు పరీక్ష", content_color="red")
    add_banner(t7, "04-09-2026 : SRI KRISHNA JANMASHTAMI", "holiday_red")
    add_banner(t7, "05-09-2026 : TEACHER’S DAY (WORKING DAY)", "event_green")
    add_banner(t7, "07-09-2026 TO 12-09-2026", "date")
    r39 = add_data_row(t7, period="1", content="పునశ్చరణ", content_color="red")
    r40 = add_data_row(t7, period="2")
    r41 = add_data_row(t7, period="3")
    r42 = add_data_row(t7, period="4")
    r39.cells[2].merge(r40.cells[2]).merge(r41.cells[2]).merge(r42.cells[2])
    add_banner(t7, "12-09-2026 : SECOND SATURDAY", "holiday_red")
    
    # Page 8
    print("[8/8] Building Page 8...")
    add_page_break_zero(doc)
    add_header_bar(doc)
    t8 = init_schedule_table(doc)
    add_banner(t8, "14-09-2026 TO 19-09-2026", "date")
    r43 = add_data_row(t8, period="1", content="పునశ్చరణ", content_color="red")
    r44 = add_data_row(t8, period="2")
    r45 = add_data_row(t8, period="3")
    r43.cells[2].merge(r44.cells[2]).merge(r45.cells[2])
    add_banner(t8, "14-09-2026 : GANESH CHATURTHI", "holiday_red")
    add_banner(t8, "19-09-2026 : NO BAG DAY-STEAM&CCA", "event_purple")
    add_banner(t8, "21-09-2026 TO 25-09-2026", "date")
    r46 = add_data_row(t8, period="1", content="పునశ్చరణ", content_color="red")
    r47 = add_data_row(t8, period="2")
    r48 = add_data_row(t8, period="3")
    r49 = add_data_row(t8, period="4")
    r46.cells[2].merge(r47.cells[2]).merge(r48.cells[2]).merge(r49.cells[2])
    add_banner(t8, "22-09-2026 TO 24-09-2026 : PTM", "event_purple")
    add_banner(t8, "26-09-2026 TO 29-09-2026 : MID TERM REVISION", "blue_banner")
    add_banner(t8, "30-09-2026 TO 09-10-2026 : MID TERM EXAMS", "blue_banner")
    add_banner(t8, "02-10-2026 : GANDHI JAYANTI", "holiday_red")
    add_banner(t8, "10-10-2026 : SECOND SATURDAY", "holiday_red")
    add_banner(t8, "12-10-2026 TO 21-10-2026 : DUSSEHRA VACATION", "holiday_red")
    
    # Save files
    local_output = "GRADE_II_TELUGU_TL_TERM-1_MCS_AP_2026-27.docx"
    doc.save(local_output)
    print(f"[OK] Saved document in current folder: {os.path.abspath(local_output)}")
    
    downloads_output = os.path.expanduser(r"~\Downloads\GRADE_II_TELUGU_TL_TERM-1_MCS_AP_2026-27.docx")
    try:
        doc.save(downloads_output)
        print(f"[OK] Saved copy to Downloads: {downloads_output}")
    except Exception as e:
        print(f"[Note] Could not save to Downloads ({e})")

    print("\nSUCCESS! Conversion completed without any language or font issues.")

if __name__ == "__main__":
    convert()
