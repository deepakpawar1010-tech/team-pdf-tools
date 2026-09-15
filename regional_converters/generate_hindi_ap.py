import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=45, bottom=45, left=65, right=65):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def format_text_run(run, text, font_name="Nirmala UI", size_pt=9.5, bold=False, color_rgb=None, italic=False):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def add_header_bar(doc, term_text="AY-2026-27 (Term - 2)"):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    col_widths = [Inches(3.2), Inches(4.3), Inches(3.2)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
    
    # Left pill: Micro schedule
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "DC2626")
    set_cell_margins(c0, top=50, bottom=50, left=140, right=140)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run()
    format_text_run(r0, "Micro schedule", font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    # Center pill: Term
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "EA580C")
    set_cell_margins(c1, top=50, bottom=50, left=140, right=140)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run()
    format_text_run(r1, term_text, font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))
    
    # Right pill: State
    c2 = tbl.rows[0].cells[2]
    set_cell_background(c2, "DC2626")
    set_cell_margins(c2, top=50, bottom=50, left=140, right=140)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run()
    format_text_run(r2, "Andhra Pradesh", font_name="Arial", size_pt=13, bold=True, color_rgb=RGBColor(255, 255, 255))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

def add_info_box_hindi(doc):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="CBD5E1", sz="6", val="single")
    
    col_widths = [Inches(2.5), Inches(2.7), Inches(2.3), Inches(3.2)]
    for i, w in enumerate(col_widths):
        tbl.rows[0].cells[i].width = w
    
    c0 = tbl.rows[0].cells[0]
    set_cell_background(c0, "FEF2F2")
    set_cell_margins(c0, top=50, bottom=50, left=100, right=100)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(0)
    r0_1 = p0.add_run("No.of Periods\n")
    format_text_run(r0_1, "No.of Periods\n", font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(185, 28, 28))
    r0_2 = p0.add_run("Per Week: 5")
    format_text_run(r0_2, "Per Week: 5", font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(185, 28, 28))
    
    c1 = tbl.rows[0].cells[1]
    set_cell_background(c1, "FFFBEB")
    set_cell_margins(c1, top=50, bottom=50, left=100, right=100)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1_1 = p1.add_run("October to March\n")
    format_text_run(r1_1, "October to March\n", font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(180, 83, 9))
    r1_2 = p1.add_run("2026 - 27")
    format_text_run(r1_2, "2026 - 27", font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(180, 83, 9))
    
    c2 = tbl.rows[0].cells[2]
    set_cell_background(c2, "EFF6FF")
    set_cell_margins(c2, top=50, bottom=50, left=100, right=100)
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Grade: II")
    format_text_run(r2, "Grade: II", font_name="Calibri", size_pt=12, bold=True, color_rgb=RGBColor(29, 78, 216))
    
    c3 = tbl.rows[0].cells[3]
    set_cell_background(c3, "F5F3FF")
    set_cell_margins(c3, top=50, bottom=50, left=100, right=100)
    c3.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p3 = c3.paragraphs[0]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(0)
    r3_1 = p3.add_run("HINDI\n")
    format_text_run(r3_1, "HINDI\n", font_name="Calibri", size_pt=11, bold=True, color_rgb=RGBColor(67, 56, 202))
    r3_2 = p3.add_run("Subject: SECOND LANGUAGE")
    format_text_run(r3_2, "Subject: SECOND LANGUAGE", font_name="Calibri", size_pt=10.5, bold=True, color_rgb=RGBColor(185, 28, 28))

    gap = doc.add_paragraph()
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.space_after = Pt(2)
    gap.paragraph_format.line_spacing = Pt(1)

def add_page_break_zero(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

COL_WIDTHS_5 = [Inches(1.0), Inches(2.2), Inches(4.2), Inches(1.8), Inches(1.5)]

def init_schedule_table(doc):
    tbl = doc.add_table(rows=0, cols=5)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl, color="94A3B8", sz="4", val="single")
    
    row = tbl.add_row()
    set_row_cant_split(row)
    headers = ["PERIOD", "CHAPTER/LESSON", "CONTENT TO BE TAUGHT", "TLM", "PRACTICE WORK"]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        cell.width = COL_WIDTHS_5[i]
        set_cell_background(cell, "1E3A8A") # Navy Blue
        set_cell_margins(cell, top=65, bottom=65, left=65, right=65)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        format_text_run(r, h, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(255, 255, 255))
    return tbl

def add_banner(tbl, text, b_type="date"):
    row = tbl.add_row()
    set_row_cant_split(row)
    cell = row.cells[0]
    for c in row.cells[1:]:
        cell.merge(c)
    
    set_cell_margins(cell, top=45, bottom=45, left=80, right=80)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    if b_type == "date":
        set_cell_background(cell, "BFDBFE")
        r = p.add_run(text)
        format_text_run(r, text, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(15, 23, 42))
    elif b_type == "holiday_red":
        set_cell_background(cell, "FEE2E2")
        r = p.add_run(text)
        format_text_run(r, text, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(185, 28, 28))
    elif b_type == "event_purple":
        set_cell_background(cell, "F3E8FF")
        r = p.add_run(text)
        format_text_run(r, text, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(107, 33, 168))
    elif b_type == "blue_banner":
        set_cell_background(cell, "E0E7FF")
        r = p.add_run(text)
        format_text_run(r, text, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(30, 58, 138))
    return row

def add_data_row(tbl, period="", chapter="", content="", tlm="", practice="", 
                 chapter_color="green", content_color="black"):
    row = tbl.add_row()
    set_row_cant_split(row)
    vals = [str(period), str(chapter), str(content), str(tlm), str(practice)]
    
    for i in range(5):
        cell = row.cells[i]
        cell.width = COL_WIDTHS_5[i]
        set_cell_margins(cell, top=40, bottom=40, left=55, right=55)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        val = vals[i]
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                r = p.add_run(val)
                format_text_run(r, val, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(15, 23, 42))
        elif i == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                r = p.add_run(val)
                col = RGBColor(16, 124, 65) if chapter_color == "green" else RGBColor(185, 28, 28)
                format_text_run(r, val, font_name="Nirmala UI", size_pt=9.5, bold=True, color_rgb=col)
        elif i == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                r = p.add_run(val)
                if content_color == "red" or "लघु-परीक्षा" in val:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    format_text_run(r, val, font_name="Nirmala UI", size_pt=10, bold=True, color_rgb=RGBColor(185, 28, 28))
                else:
                    format_text_run(r, val, font_name="Nirmala UI", size_pt=9, bold=False, color_rgb=RGBColor(15, 23, 42))
        elif i == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if val:
                lines = val.split("\\n")
                for li, l in enumerate(lines):
                    if li > 0:
                        p.add_run("\\n")
                    r = p.add_run(l)
                    if "A.V." in l:
                        format_text_run(r, l, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(15, 23, 42))
                    else:
                        format_text_run(r, l, font_name="Nirmala UI", size_pt=8.5, bold=True, color_rgb=RGBColor(15, 23, 42))
        elif i == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if val:
                r = p.add_run(val)
                format_text_run(r, val, font_name="Nirmala UI", size_pt=8.5, bold=False, color_rgb=RGBColor(30, 41, 59))
    return row

def generate_hindi_doc(output_path):
    print("Generating Hindi Micro Schedule Word Document (Grade II Term 2)...")
    doc = docx.Document()
    for s in doc.sections:
        s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        s.page_width = Inches(11.69)
        s.page_height = Inches(8.27)
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    # ==================== PAGE 1 ====================
    print("  -> Page 1...")
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    add_info_box_hindi(doc)
    t1 = init_schedule_table(doc)
    
    add_banner(t1, "22-10-2026 TO 24-10-2026", "date")
    r1 = add_data_row(t1, period="1", chapter="‘ए’ की मात्रा", 
                      content="पाठ्य पुस्तक पृष्ठ संख्या-1 का अभ्यास", 
                      tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r2 = add_data_row(t1, period="2", chapter="पाठ-7 किसके पीछे (ए)", 
                      content="पाठ्य पुस्तक पृष्ठ संख्या-2-3 का अभ्यास")
    r1.cells[3].merge(r2.cells[3])
    
    add_banner(t1, "24-10-2026: NO BAG DAY - STEAM & CCA", "event_purple")
    add_banner(t1, "26-10-2026 TO 31-10-2026: SLC-1 / BUFFER", "event_purple")
    add_banner(t1, "02-11-2026 TO 07-11-2026", "date")
    
    r3 = add_data_row(t1, period="1", chapter="पाठ-7 किसके पीछे (ए)", 
                      content="पाठ्य पुस्तक पृष्ठ संख्या-4-5 का अभ्यास")
    r4 = add_data_row(t1, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-6 का अभ्यास")
    r5 = add_data_row(t1, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-7 का अभ्यास")
    r6 = add_data_row(t1, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-8-10 का अभ्यास")
    r3.cells[1].merge(r4.cells[1]).merge(r5.cells[1]).merge(r6.cells[1])
    
    add_data_row(t1, period="5", chapter="पाठ-7 किसके पीछे (ए)", 
                 content="लघु-परीक्षा", chapter_color="red", content_color="red")
                 
    add_banner(t1, "05-11-2026 & 06-11-2026: PTM", "event_purple")
    add_banner(t1, "07-11-2026: PTM/ NO BAG DAY / CCA / CLUB", "event_purple")
    add_banner(t1, "09-11-2026 TO 14-11-2026", "date")
    
    add_data_row(t1, period="1", chapter="‘ऐ’ की मात्रा\\nपाठ-8 रेत का घर (ऐ)", 
                 content="पाठ्य पुस्तक पृष्ठ संख्या-11 का अभ्यास", 
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")

    # ==================== PAGE 2 ====================
    print("  -> Page 2...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t2 = init_schedule_table(doc)
    
    r7 = add_data_row(t2, period="2", chapter="पाठ-8 रेत का घर (ऐ)", content="पाठ्य पुस्तक पृष्ठ संख्या-12 का अभ्यास")
    r8 = add_data_row(t2, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-13 का अभ्यास")
    r9 = add_data_row(t2, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-14 का अभ्यास")
    r7.cells[1].merge(r8.cells[1]).merge(r9.cells[1])
    
    add_banner(t2, "13-11-2026: BUFFER", "event_purple")
    add_banner(t2, "14-11-2026: CHILDREN'S DAY / SECOND SATURDAY", "holiday_red")
    add_banner(t2, "16-11-2026 TO 21-11-2026", "date")
    
    r10 = add_data_row(t2, period="1", chapter="पाठ-8 रेत का घर (ऐ)", content="पाठ्य पुस्तक पृष्ठ संख्या-15-16 का अभ्यास")
    r11 = add_data_row(t2, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-17 का अभ्यास")
    r12 = add_data_row(t2, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-18 का अभ्यास")
    r13 = add_data_row(t2, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-19-20 का अभ्यास")
    r10.cells[1].merge(r11.cells[1]).merge(r12.cells[1]).merge(r13.cells[1])
    
    add_data_row(t2, period="5", chapter="पाठ-8 रेत का घर (ऐ)", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    add_banner(t2, "21-11-2026: NO BAG DAY-STEAM & CCA", "event_purple")
    add_banner(t2, "23-11-2026 TO 28-11-2026", "date")
    
    r14 = add_data_row(t2, period="1", chapter="‘ओ’ की मात्रा", content="पाठ्य पुस्तक पृष्ठ संख्या-21 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r15 = add_data_row(t2, period="2", chapter="पाठ-9 जोकर का खेल (ओ)", content="पाठ्य पुस्तक पृष्ठ संख्या-22 का अभ्यास")
    r14.cells[3].merge(r15.cells[3])
    
    r16 = add_data_row(t2, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-23 का अभ्यास")
    r17 = add_data_row(t2, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-24 का अभ्यास")
    r18 = add_data_row(t2, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-25 का अभ्यास")
    r15.cells[1].merge(r16.cells[1]).merge(r17.cells[1]).merge(r18.cells[1])
    
    add_banner(t2, "24-11-2026: GURUNANAK JAYANTHI", "holiday_red")

    # ==================== PAGE 3 ====================
    print("  -> Page 3...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t3 = init_schedule_table(doc)
    
    add_banner(t3, "30-11-2026 TO 05-12-2026", "date")
    r19 = add_data_row(t3, period="1", chapter="पाठ-9 जोकर का खेल (ओ)", content="पाठ्य पुस्तक पृष्ठ संख्या-26-27 का अभ्यास")
    r20 = add_data_row(t3, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-28-30 का अभ्यास")
    r19.cells[1].merge(r20.cells[1])
    
    add_data_row(t3, period="3", chapter="‘औ’ की मात्रा", content="पाठ्य पुस्तक पृष्ठ संख्या-31 का अभ्यास")
    
    r21 = add_data_row(t3, period="4", chapter="पाठ-\\n10 मटर या टमाटर (औ)", content="पाठ्य पुस्तक पृष्ठ संख्या-32 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r22 = add_data_row(t3, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-33 का अभ्यास")
    r21.cells[1].merge(r22.cells[1])
    
    add_banner(t3, "05-12-2026: NO BAG DAY/BB/CCA", "event_purple")
    add_banner(t3, "07-12-2026 TO 12-12-2026", "date")
    
    r23 = add_data_row(t3, period="1", chapter="पाठ-\\n10 मटर या टमाटर (औ)", content="पाठ्य पुस्तक पृष्ठ संख्या-34 का अभ्यास")
    r24 = add_data_row(t3, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-35-36 का अभ्यास")
    r25 = add_data_row(t3, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-37-38 का अभ्यास")
    r26 = add_data_row(t3, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-39-41 का अभ्यास")
    r23.cells[1].merge(r24.cells[1]).merge(r25.cells[1]).merge(r26.cells[1])
    
    add_data_row(t3, period="5", chapter="पाठ-\\n10 मटर या टमाटर (औ)", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    add_banner(t3, "12-12-2026: SECOND SATURDAY", "holiday_red")
    add_banner(t3, "14-12-2026 TO 19-12-2026", "date")
    
    add_data_row(t3, period="1", chapter="‘अं’, ‘अः’, ‘अँ’ की मात्रा", content="पाठ्य पुस्तक पृष्ठ संख्या-42 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r27 = add_data_row(t3, period="2", chapter="पाठ-11 अंगूर किसने खाए", content="पाठ्य पुस्तक पृष्ठ संख्या-43 का अभ्यास")
    r28 = add_data_row(t3, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-44 का अभ्यास")
    r29 = add_data_row(t3, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-45 का अभ्यास")
    r30 = add_data_row(t3, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-46 का अभ्यास")
    r27.cells[1].merge(r28.cells[1]).merge(r29.cells[1]).merge(r30.cells[1])
    
    add_banner(t3, "19-12-2026: NO BAG DAY-STEAM&CCA", "event_purple")

    # ==================== PAGE 4 ====================
    print("  -> Page 4...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t4 = init_schedule_table(doc)
    
    add_banner(t4, "21-12-2026 TO 26-12-2026", "date")
    r31 = add_data_row(t4, period="1", chapter="पाठ-11 अंगूर किसने खाए", content="पाठ्य पुस्तक पृष्ठ संख्या-47 का अभ्यास")
    r32 = add_data_row(t4, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-48-50 का अभ्यास")
    r31.cells[1].merge(r32.cells[1])
    add_data_row(t4, period="3", chapter="पाठ-11 अंगूर किसने खाए", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r33 = add_data_row(t4, period="4", chapter="बारहखड़ी", content="पाठ्य पुस्तक पृष्ठ संख्या-51 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r34 = add_data_row(t4, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-52 का अभ्यास")
    r33.cells[1].merge(r34.cells[1])
    
    add_banner(t4, "21-12-2026 & 22-12-2026: PTM", "event_purple")
    add_banner(t4, "25-12-2026: CHRISTMAS", "holiday_red")
    add_banner(t4, "28-12-2026 TO 02-01-2027", "date")
    
    add_data_row(t4, period="1", chapter="बारहखड़ी", content="पाठ्य पुस्तक पृष्ठ संख्या-53-55 का अभ्यास")
    add_data_row(t4, period="2", chapter="बारहखड़ी", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    add_data_row(t4, period="3", chapter="सोचो और बताओ", content="पाठ्य पुस्तक पृष्ठ संख्या-56 का अभ्यास")
    add_data_row(t4, period="4", chapter="वर्ग-पहेली, भाषा-खेल", content="पाठ्य पुस्तक पृष्ठ संख्या-57-58 का अभ्यास")
    add_data_row(t4, period="5", chapter="द्वित्वाक्षर", content="पाठ्य पुस्तक पृष्ठ संख्या-59 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
                 
    add_banner(t4, "01-01-2027: NEW YEAR", "holiday_red")
    add_banner(t4, "04-01-2027 TO 09-01-2027", "date")
    
    add_data_row(t4, period="1", chapter="द्वित्वाक्षर", content="पाठ्य पुस्तक पृष्ठ संख्या-60-61 का अभ्यास")
    add_data_row(t4, period="2", chapter="द्वित्वाक्षर", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r35 = add_data_row(t4, period="3", chapter="संयुक्ताक्षर", content="पाठ्य पुस्तक पृष्ठ संख्या-62 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r36 = add_data_row(t4, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-63 का अभ्यास")
    r37 = add_data_row(t4, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-64 का अभ्यास")
    r35.cells[1].merge(r36.cells[1]).merge(r37.cells[1])
    
    add_banner(t4, "09-01-2027: SECOND SATURDAY", "holiday_red")
    add_banner(t4, "11-01-2027 TO 16-01-2027: PONGAL VACATION", "holiday_red")

    # ==================== PAGE 5 ====================
    print("  -> Page 5...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t5 = init_schedule_table(doc)
    
    add_banner(t5, "18-01-2027 TO 23-01-2027", "date")
    add_data_row(t5, period="1", chapter="संयुक्ताक्षर", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r38 = add_data_row(t5, period="2", chapter="‘र’ के रूप", content="पाठ्य पुस्तक पृष्ठ संख्या-65 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r39 = add_data_row(t5, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-66 का अभ्यास")
    r40 = add_data_row(t5, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-67 का अभ्यास")
    r41 = add_data_row(t5, period="5", content="पाठ्य पुस्तक पृष्ठ संख्या-68-69 का अभ्यास")
    r38.cells[1].merge(r39.cells[1]).merge(r40.cells[1]).merge(r41.cells[1])
    
    add_banner(t5, "23-01-2027 : NO BAG DAY-STEAM&CCA", "event_purple")
    add_banner(t5, "25-01-2027 TO 30-01-2027", "date")
    
    add_data_row(t5, period="1", chapter="‘र’ के रूप", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r42 = add_data_row(t5, period="2", chapter="वाक्य-रचना", content="पाठ्य पुस्तक पृष्ठ संख्या-70 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r43 = add_data_row(t5, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-71 का अभ्यास")
    r44 = add_data_row(t5, period="4", content="पाठ्य पुस्तक पृष्ठ संख्या-72 का अभ्यास")
    r42.cells[1].merge(r43.cells[1]).merge(r44.cells[1])
    
    add_banner(t5, "26-01-2027: REPUBLIC DAY", "holiday_red")
    add_banner(t5, "27-01-2027 TO 29-01-2027: PTM", "event_purple")
    add_banner(t5, "30-01-2027: NO BAG DAY/BB/CCA/CLUB", "event_purple")
    add_banner(t5, "01-02-2027 TO 06-02-2027", "date")
    
    r45 = add_data_row(t5, period="1", chapter="वाक्य-रचना", content="पाठ्य पुस्तक पृष्ठ संख्या-73 का अभ्यास")
    r46 = add_data_row(t5, period="2", content="पाठ्य पुस्तक पृष्ठ संख्या-74 का अभ्यास")
    r47 = add_data_row(t5, period="3", content="पाठ्य पुस्तक पृष्ठ संख्या-75-76 का अभ्यास")
    r45.cells[1].merge(r46.cells[1]).merge(r47.cells[1])
    
    r48 = add_data_row(t5, period="4", chapter="गिनती", content="पाठ्य पुस्तक पृष्ठ संख्या-77 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r49 = add_data_row(t5, period="5", content="लिखित अभ्यास करवाएँ")
    r48.cells[1].merge(r49.cells[1])
    
    add_banner(t5, "06-02-2027 : NO BAG DAY/BB/CCA", "event_purple")

    # ==================== PAGE 6 ====================
    print("  -> Page 6...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t6 = init_schedule_table(doc)
    
    add_banner(t6, "08-02-2027 TO 13-02-2027", "date")
    add_data_row(t6, period="1", chapter="गिनती", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r50 = add_data_row(t6, period="2", chapter="विलोम शब्द", content="पाठ्य पुस्तक पृष्ठ संख्या-78 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r51 = add_data_row(t6, period="3", content="मौखिक अभ्यास करवाएँ")
    r52 = add_data_row(t6, period="4", content="लिखित अभ्यास करवाएँ")
    r50.cells[1].merge(r51.cells[1]).merge(r52.cells[1])
    
    add_data_row(t6, period="5", chapter="एक-अनेक", content="पाठ्य पुस्तक पृष्ठ संख्या-79 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
                 
    add_banner(t6, "13-02-2027 : SECOND SATURDAY", "holiday_red")
    add_banner(t6, "15-02-2027 TO 20-02-2027", "date")
    
    add_data_row(t6, period="1", chapter="एक-अनेक", content="लिखित अभ्यास करवाएँ")
    add_data_row(t6, period="2", chapter="एक-अनेक, विलोम शब्द", content="लघु-परीक्षा", chapter_color="red", content_color="red")
    
    r53 = add_data_row(t6, period="3", chapter="लड़का-लड़की", content="पाठ्य पुस्तक पृष्ठ संख्या-80 का अभ्यास",
                       tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    r54 = add_data_row(t6, period="4", content="लिखित अभ्यास करवाएँ")
    r53.cells[1].merge(r54.cells[1])
    
    add_data_row(t6, period="5", chapter="नाम बताने वाले शब्द", content="पाठ्य पुस्तक पृष्ठ संख्या-81 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
                 
    add_banner(t6, "20-02-2027: NO BAG DAY-STEAM & CCA", "event_purple")
    add_banner(t6, "22-02-2027 TO 27-02-2027: SLC - 2/ BUFFER", "event_purple")
    add_banner(t6, "01-03-2027 TO 06-03-2027", "date")
    
    r55 = add_data_row(t6, period="1", chapter="नाम बताने वाले शब्द", content="मौखिक अभ्यास करवाएँ")
    r56 = add_data_row(t6, period="2", content="लिखित अभ्यास करवाएँ")
    r55.cells[1].merge(r56.cells[1])
    
    add_data_row(t6, period="3", chapter="नाम बताने वाले शब्द,\\nलड़का-लड़की", content="लघु-परीक्षा", chapter_color="red", content_color="red")

    # ==================== PAGE 7 ====================
    print("  -> Page 7...")
    add_page_break_zero(doc)
    add_header_bar(doc, "AY-2026-27 (Term - 2)")
    t7 = init_schedule_table(doc)
    
    add_data_row(t7, period="4", chapter="Reading for Fun\\nऊँट, बोली के गुण", content="पाठ्य पुस्तक पृष्ठ संख्या-82-83 का अभ्यास")
    add_data_row(t7, period="5", chapter="Reading for Fun\\nप्यासे कौए", content="पाठ्य पुस्तक पृष्ठ संख्या-84-85 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
                 
    add_banner(t7, "06-03-2027: SHIVARATHRI", "holiday_red")
    add_banner(t7, "08-03-2027 TO 10-03-2027", "date")
    
    add_data_row(t7, period="1", chapter="Reading for Fun\\nकछुआ और खरगोश", content="पाठ्य पुस्तक पृष्ठ संख्या-86-87 का अभ्यास",
                 tlm="A.V. Modules\\nश्रव्य और दृश्य संसाधन")
    add_data_row(t7, period="2", chapter="Reading for Fun\\nआओ कहानी बनाएँ,\\nखेल-खेल में", content="पाठ्य पुस्तक पृष्ठ संख्या-88-90 का अभ्यास")
    
    add_banner(t7, "10-03-2027: RAMZAN", "holiday_red")
    add_banner(t7, "11-03-2027 TO 15-03-2027: TERM-END REVISION", "blue_banner")
    add_banner(t7, "13-03-2027: SECOND SATURDAY", "holiday_red")
    add_banner(t7, "16-03-2027 TO 25-03-2027: TERM-END EXAMS", "blue_banner")
    add_banner(t7, "22-03-2027: HOLI", "holiday_red")
    add_banner(t7, "26-03-2027: GOOD FRIDAY", "holiday_red")
    add_banner(t7, "27-03-2027 TO 31-03-2027: SHORT VACATION", "holiday_red")

    doc.save(output_path)
    print(f"Saved: {output_path}")

if __name__ == "__main__":
    out_file = r"C:\Users\Deepak Pawar\Desktop\Automation Projects\Telugu_PDF_to_Word\output_docs\GRADE-2_SOUTH HINDI (SL)_AP_MCS_TERM 2_AY 2026-27.docx"
    generate_hindi_doc(out_file)
